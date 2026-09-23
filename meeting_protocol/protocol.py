"""Проверяемый протокол и экспорт, не зависящие от GPU и моделей."""
import hashlib
import json
from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt
from pydantic import BaseModel, Field, model_validator

ROOT = Path(__file__).resolve().parent.parent


class Segment(BaseModel):
    id: int
    start: float = Field(ge=0)
    end: float = Field(ge=0)
    text: str
    speaker: str | None = None


class Assignment(BaseModel):
    task: str = Field(min_length=1)
    responsible: str | None = Field(description="Ответственный человек или подразделение из транскрипта; null только если не указан")
    deadline_text: str | None = Field(description="Окончательный согласованный срок исходными словами; null только если не указан")
    evidence: list[int] = Field(min_length=1)


class SpeakerName(BaseModel):
    label: str = Field(description="Метка диаризации, например SPEAKER_2")
    name: str | None = Field(description="Имя из списка участников; null, если в транскрипте нет явного признака")
    evidence: list[int] = Field(default_factory=list, description="Номера реплик, из которых следует имя")


class Extraction(BaseModel):
    assignments: list[Assignment]
    summary: str
    speakers: list[SpeakerName] = Field(default_factory=list)


class MeetingContext(BaseModel):
    participants: list[str] = Field(default_factory=list, max_length=100)
    topic: str = Field(default="", max_length=1000)

    @model_validator(mode="after")
    def clean(self):
        self.participants = list(dict.fromkeys(name.strip() for name in self.participants if name.strip()))
        if any(len(name) > 200 for name in self.participants):
            raise ValueError("Имя участника длиннее 200 символов")
        self.topic = self.topic.strip()
        return self


class Protocol(Extraction):
    source_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    source_name: str
    meeting_date: date
    segments: list[Segment] = Field(min_length=1)
    provenance: dict[str, str]
    context: MeetingContext = Field(default_factory=MeetingContext)

    @model_validator(mode="after")
    def validate_evidence(self):
        ids = {segment.id for segment in self.segments}
        if len(ids) != len(self.segments):
            raise ValueError("Идентификаторы реплик должны быть уникальными")
        for assignment in self.assignments:
            if not set(assignment.evidence) <= ids:
                raise ValueError("Поручение ссылается на отсутствующую реплику")
        labels = {segment.speaker for segment in self.segments}
        if any(speaker.label not in labels or not set(speaker.evidence) <= ids for speaker in self.speakers):
            raise ValueError("Имя говорящего ссылается на отсутствующую метку или реплику")
        return self

    def speaker_title(self, label: str | None) -> str:
        if not label:
            return "Говорящий не определён"
        name = next((s.name for s in self.speakers if s.label == label and s.name), None)
        return f"{label} ({name})" if name else label


def audio_hash(path: Path) -> str:
    with path.open("rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest()


def meeting_presets() -> dict:
    return json.loads((ROOT / "data" / "meeting-context.json").read_text(encoding="utf-8"))


def resolve_context(audio: Path, participants: str | None = None, topic: str | None = None) -> MeetingContext:
    preset = MeetingContext.model_validate(meeting_presets().get(audio_hash(audio), {}))
    return MeetingContext(
        participants=preset.participants if participants is None else participants.replace(";", "\n").splitlines(),
        topic=preset.topic if topic is None else topic,
    )


def replay(audio: Path, *, context: MeetingContext | None = None) -> Protocol:
    digest = audio_hash(audio)
    for path in sorted((ROOT / "data" / "replay").glob("*.json")):
        protocol = Protocol.model_validate_json(path.read_text())
        if protocol.source_sha256 == digest:
            if context is not None and context != protocol.context:
                raise ValueError("Участники или тема отличаются от сохранённых: воспроизведение не пересчитывает протокол; нужен режим GPU")
            return protocol
    raise ValueError("Нет сохранённого результата для этого аудио. Нужен режим GPU.")


def export_docx(protocol: Protocol, *, replay_mode: bool = False) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Протокол совещания", 0)
    doc.add_paragraph(f"Запись: {protocol.source_name}. Дата встречи: {protocol.meeting_date}.")
    if protocol.context.topic:
        doc.add_paragraph(f"Тема: {protocol.context.topic}")
    if protocol.context.participants:
        doc.add_paragraph("Участники и упомянутые ответственные: " + "; ".join(protocol.context.participants))
    if replay_mode:
        doc.add_paragraph("Воспроизведение заранее вычисленного результата — не проверка моделей.")
    doc.add_paragraph("Черновик ИИ: проверьте имена, сроки и содержание по записи.")
    if protocol.speakers:
        doc.add_paragraph("Говорящие: " + "; ".join(protocol.speaker_title(s.label) for s in protocol.speakers)
                          + ". Метки получены автоматически по голосу, имена — по обращениям в разговоре.")
    doc.add_heading("Саммари", 1)
    doc.add_paragraph(protocol.summary or "Не сформировано")
    doc.add_heading("Поручения", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ["Суть", "Ответственный", "Срок (как сказано)", "Реплики"]):
        cell.text = title
    for assignment in protocol.assignments:
        values = [assignment.task, assignment.responsible or "Не указан", assignment.deadline_text or "Не указан", ", ".join(map(str, assignment.evidence))]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    doc.add_heading("Транскрипт", 1)
    for segment in protocol.segments:
        doc.add_paragraph(f"[{segment.id}; {segment.start:.1f}–{segment.end:.1f} с] {protocol.speaker_title(segment.speaker)}: {segment.text}")
    doc.add_heading("Происхождение", 1)
    doc.add_paragraph(f"SHA-256: {protocol.source_sha256}")
    for key, value in protocol.provenance.items():
        doc.add_paragraph(f"{key}: {value}")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
