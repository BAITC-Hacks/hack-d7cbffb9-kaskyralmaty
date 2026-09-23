"""Проверка артефактов через их публичные форматы."""
import hashlib
import json
from pathlib import Path
from docx import Document

for number in (1, 2):
    prefix = Path(f"outputs/meeting-{number}")
    result = json.loads(prefix.with_suffix(".json").read_text())
    source = Path(f"docs/Трек 8 Инновации/Совещание №{number}.mp3")
    assert result["source_sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    preset = json.loads(Path('data/meeting-context.json').read_text())[result['source_sha256']]
    context = {key: preset[key] for key in ('participants', 'topic')}
    assert result["context"] == context, "Не сохранены тема и список участников"
    allowed = set(context['participants']) | {"Юридический департамент", "юридический департамент"}
    assert all(a["responsible"] is None or a["responsible"] in allowed for a in result["assignments"]), "Искажённое имя в поручениях комплектного примера"
    assert len(result["segments"]) > 5, "Отсутствует транскрипт всей встречи"
    assert result["segments"][-1]["end"] > 150, "Протокол покрывает лишь начало встречи"
    assert result["assignments"], "Поручения не извлечены"
    assert sum(bool(a["responsible"]) for a in result["assignments"]) >= 3, "Пропущены явно названные исполнители"
    assert sum(bool(a["deadline_text"]) for a in result["assignments"]) >= 3, "Пропущены явно названные сроки"
    doc = Document(prefix.with_suffix(".docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert result["segments"][0]["text"] in text
    assert context['topic'] in text
    assert len(doc.tables[0].rows) == len(result["assignments"]) + 1
    print(f"Документ №{number}: {len(result['segments'])} реплик, {len(result['assignments'])} поручений")
