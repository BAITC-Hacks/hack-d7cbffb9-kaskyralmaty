"""Проверяем публичный результат: документ и безопасное воспроизведение."""
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from meeting_protocol.protocol import Assignment, Protocol, Segment, export_docx, replay


def test_docx_preserves_kazakh_and_assignment():
    protocol = Protocol(
        source_sha256="a" * 64, source_name="recording.mp3", meeting_date=date(2026, 9, 23),
        segments=[Segment(id=0, start=0, end=3, text="Динара, жұмаға дейін дайындаңыз.")],
        assignments=[Assignment(task="Подготовить договор", responsible="Динара", deadline_text="жұмаға дейін", evidence=[0])],
        summary="Обсудили договор.", provenance={"asr": "test"},
    )
    doc = Document(BytesIO(export_docx(protocol)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "жұмаға дейін" in text
    assert "Обсудили договор" in text
    assert "Динара" in doc.tables[0].cell(1, 1).text


def test_replay_rejects_unknown_audio(tmp_path):
    audio = tmp_path / "unknown.mp3"
    audio.write_bytes(b"not a known recording")
    with pytest.raises(ValueError, match="Нет сохранённого результата"):
        replay(audio)
